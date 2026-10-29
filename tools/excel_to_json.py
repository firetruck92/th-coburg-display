#!/usr/bin/env python3
"""Wandelt die Excel-Zeitplanung (Hochschule Coburg, Layout "Zeitplanung_WI") in schedule.json und schedule.js um.

Aufruf:  python3 excel_to_json.py Zeitplanung.xlsx [ausgabe-ordner]

Was gelesen wird (erstes sichtbares Tabellenblatt):
  - Modulblöcke: Kopfzeile mit Modulname und (Dozent) je Semester, darunter eine Zeile pro Datum
  - Spalten E/F = 1. Semester, H/I = 3. Semester, K.. = 5. und 7. Semester (Zeiten / Ort)
  - Spalten Q:S = Studiengangsrelevante Termine, Spalte U = Allgemeines (Feiertage, Ferien, Aktionen)
  - rote Schrift = noch vorläufig
Teams-Links werden bewusst NICHT übernommen, nur das Wort "Online".
"""
import sys, re, json, datetime as dt, pathlib
import openpyxl

src = pathlib.Path(sys.argv[1])
outdir = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else pathlib.Path('.')
outdir.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(src)
ws = next(s for s in wb.worksheets if s.sheet_state == 'visible')

warnings = []

def is_red(c):
    col = c.font.color
    return bool(col is not None and col.type == 'rgb' and col.rgb in ('FFFF0000', 'FFFF0000'))

def clean(s):
    return re.sub(r'\s+', ' ', str(s).replace('\t', ' ')).strip()

def merge_of(cell):
    for m in ws.merged_cells.ranges:
        if cell.coordinate in m:
            return m
    return None

TIME_RE = re.compile(r'(\d{1,2})[.:](\d{2})\s*[-–]\s*(\d{1,2})[.:](\d{2})')
def fmt(h, m): return f'{int(h):02d}:{int(m):02d}'
def parse_range(v):
    if v is None: return None
    m = TIME_RE.search(str(v))
    return (fmt(m[1], m[2]), fmt(m[3], m[4])) if m else None

def parse_room(v):
    """-> (room|None, online, tbd)"""
    if v is None: return (None, False, False)
    s = str(v).strip()
    if s.upper().startswith('=HYPERLINK'):
        return ('Online', True, False)
    s = clean(s)
    if s.lower() in ('tba', 'tbd', ''): return (None, False, True)
    if s.lower() == 'online': return ('Online', True, False)
    return (s, False, False)

def parse_header(v):
    t = clean(v)
    groups = re.findall(r'\(([^)]*)\)', t)
    title = clean(re.sub(r'\([^)]*\)', lambda m: m[0] if 'semester' in m[0].lower() else '', t))
    who, dual = None, False
    for g in groups:
        gl = g.strip().lower()
        if gl == 'dual': dual = True
        elif 'semester' in gl: pass
        else: who = g.strip()
    return title, who, dual

def d_iso(v): return v.strftime('%Y-%m-%d')

# --- Kopfzeilen finden ---
HDR = []
for r in range(4, ws.max_row + 1):
    e = ws.cell(r, 5).value
    if ws.cell(r, 2).value is None and isinstance(e, str) and not e.startswith('Achtung'):
        HDR.append(r)
HDR.append(ws.max_row + 1)

lessons = []
for i in range(len(HDR) - 1):
    h, nxt = HDR[i], HDR[i + 1]
    # Datumszeilen dieses Blocks
    rows = [r for r in range(h + 1, nxt) if hasattr(ws.cell(r, 2).value, 'year')]
    e_cell = ws.cell(h, 5)
    em = merge_of(e_cell)
    if em and em.max_col >= 15:           # WPF-Block: gilt für alle Semester, Zeit in E, Ort in K
        title, who, dual = parse_header(e_cell.value)
        groups = [([1, 3, 5, 7], title, who, dual, 5, 11)]
    else:
        groups = []
        for col, sems in ((5, [1]), (8, [3])):
            c = ws.cell(h, col)
            if c.value:
                t, w, d = parse_header(c.value)
                groups.append((sems, t, w, d, col, col + 1))
        kc = ws.cell(h, 11)
        if kc.value:
            km = merge_of(kc)
            t, w, d = parse_header(kc.value)
            if km and km.max_col >= 15:   # K:O -> 5. und 7. Semester gemeinsam, Zeit K, Ort N
                groups.append(([5, 7], t, w, d, 11, 14))
            else:                         # nur 5. Semester, Zeit K, Ort L
                groups.append(([5], t, w, d, 11, 12))
        nc = ws.cell(h, 14)
        if nc.value:
            t, w, d = parse_header(nc.value)
            groups.append(([7], t, w, d, 14, 15))
    for r in rows:
        date = ws.cell(r, 2).value
        for sems, title, who, dual, tcol, rcol in groups:
            tc, rc = ws.cell(r, tcol), ws.cell(r, rcol)
            if tc.value is None and rc.value is None: continue
            rng = parse_range(tc.value)
            if not rng:
                warnings.append(f'Zeile {r} ({d_iso(date)}), {title}: Zeit nicht lesbar: {tc.value!r}')
                continue
            room, online, tbd = parse_room(rc.value)
            lessons.append({
                'date': d_iso(date), 'sems': sems, 'title': title, 'who': who, 'dual': dual,
                'start': rng[0], 'end': rng[1], 'room': room, 'online': online, 'tbd': tbd,
                'provisional': any(is_red(c) for c in (tc, rc) if c.value is not None)
            })

lessons.sort(key=lambda x: (x['date'], x['start'], x['sems'], x['title']))

# --- Studiengangsrelevante Termine (Q:S) ---
def time_cell(v):
    if v is None: return None
    if isinstance(v, dt.time): return v.strftime('%H:%M')
    s = str(v)
    rs = [f'{fmt(a,b)}–{fmt(c,d)}' for a, b, c, d in TIME_RE.findall(s)]
    if rs: return ', '.join(rs)
    return None if s.strip().lower() in ('tba', 'tbd') else clean(s)

events = []
for r in range(4, ws.max_row + 1):
    q = ws.cell(r, 17).value
    date = ws.cell(r, 2).value
    if q and hasattr(date, 'year'):
        room, online, tbd = parse_room(ws.cell(r, 19).value)
        title = ' / '.join(clean(x) for x in str(q).split('\n') if x.strip())
        events.append({'date': d_iso(date), 'title': title, 'time': time_cell(ws.cell(r, 18).value),
                       'place': room, 'online': online, 'tbd': tbd,
                       'provisional': any(is_red(ws.cell(r, c)) for c in (17, 18, 19) if ws.cell(r, c).value is not None)})

# --- Allgemeines (U): Feiertage, Ferien, Aktionen; aufeinanderfolgende gleiche Einträge werden zusammengefasst ---
notes = []
for r in range(4, ws.max_row + 1):
    u = ws.cell(r, 21).value
    date = ws.cell(r, 2).value
    if u and hasattr(date, 'year'):
        text = clean(u); d = d_iso(date)
        if notes and notes[-1]['text'] == text and (date.date() - dt.date.fromisoformat(notes[-1]['to'])).days <= 1:
            notes[-1]['to'] = d
        else:
            notes.append({'from': d, 'to': d, 'text': text})

meta = {'source': src.name, 'sheet': ws.title, 'generated': dt.datetime.now().isoformat(timespec='seconds'),
        'semesters': [1, 3, 5, 7]}
data = {'meta': meta, 'lessons': lessons, 'events': events, 'notes': notes}

(outdir / 'schedule.json').write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding='utf-8')
(outdir / 'schedule.js').write_text('window.SCHEDULE = ' + json.dumps(data, ensure_ascii=False) + ';\n', encoding='utf-8')

print(f'Blatt: {ws.title}')
print(f'Unterrichtseinträge: {len(lessons)}, Termine: {len(events)}, Hinweise: {len(notes)}')
for s in (1, 3, 5, 7):
    n = sum(1 for l in lessons if s in l['sems'])
    print(f'  {s}. Semester: {n} Einträge')
print('vorläufig:', [(l['date'], l['title']) for l in lessons if l['provisional']], [(e['date'], e['title']) for e in events if e['provisional']])
for w in warnings: print('WARNUNG', w)
